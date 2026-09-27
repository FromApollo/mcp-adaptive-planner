import os
from mcp.server.fastmcp.server import FastMCP

mcp = FastMCP("reports")


def _ensure_dir(path: str) -> None:
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)


@mcp.tool()
def create_pdf(title: str, content: str, output_path: str) -> str:
    """Create a PDF report with a title and multi-line text content."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    _ensure_dir(output_path)
    doc = SimpleDocTemplate(
        output_path, pagesize=A4,
        rightMargin=2 * cm, leftMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm,
    )
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 0.5 * cm)]

    for line in content.splitlines():
        stripped = line.strip()
        if stripped:
            story.append(Paragraph(stripped, styles["Normal"]))
            story.append(Spacer(1, 0.15 * cm))

    doc.build(story)
    return f"PDF created: {output_path}"


@mcp.tool()
def create_docx(title: str, content: str, output_path: str) -> str:
    """Create a DOCX report with a title and multi-line text content."""
    from docx import Document

    _ensure_dir(output_path)
    doc = Document()
    doc.add_heading(title, level=0)

    for line in content.splitlines():
        stripped = line.strip()
        if stripped:
            doc.add_paragraph(stripped)

    doc.save(output_path)
    return f"DOCX created: {output_path}"


@mcp.tool()
def create_pdf_with_table(title: str, headers: str, rows: str, output_path: str) -> str:
    """Create a PDF report with a data table.
    headers: comma-separated column names.
    rows: semicolon-separated rows, each row is comma-separated values.
    Example: headers='Name,Score' rows='Alice,95;Bob,87'
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    _ensure_dir(output_path)
    doc = SimpleDocTemplate(output_path, pagesize=A4,
                            rightMargin=2 * cm, leftMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 0.5 * cm)]

    header_row = [h.strip() for h in headers.split(",")]
    data = [header_row]
    for row_str in rows.split(";"):
        if row_str.strip():
            data.append([c.strip() for c in row_str.split(",")])

    table = Table(data)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.lightgrey]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
    ]))
    story.append(table)
    doc.build(story)
    return f"PDF with table created: {output_path}"


if __name__ == "__main__":
    mcp.run()
